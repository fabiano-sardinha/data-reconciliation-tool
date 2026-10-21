import pandas as pd
from openpyxl import load_workbook

from reconcile import run_reconciliation


DETAILS_SHEET = "Duplicate Details"

DETAILS_COLUMNS = [
    "Customer ID",
    "Source File",
    "Excel Row",
    "Name",
    "Email",
    "Status",
    "Occurrences in File",
]


def read_main(path):
    return pd.read_excel(path).set_index("Customer ID")


def read_details(path):
    return pd.read_excel(path, sheet_name=DETAILS_SHEET)


def detail_occurrences(details, customer_id):
    """Occurrences of one Customer ID as (file, excel row, email) tuples."""

    rows = details[details["Customer ID"] == customer_id]

    return [
        (row["Source File"], row["Excel Row"], row["Email"])
        for _, row in rows.iterrows()
    ]


# --- Duplicates only in File A ---

def test_duplicate_only_in_file_a_keeps_every_occurrence(
    make_excel, tmp_path
):

    file_a = make_excel("a.xlsx", [
        (1, "Ana", "ana@x.com", "Active"),
        (10, "Rui", "rui.old@x.com", "Active"),
        (10, "Rui", "rui.new@x.com", "Inactive"),
    ])
    file_b = make_excel("b.xlsx", [
        (1, "Ana", "ana@x.com", "Active"),
        (10, "Rui", "rui.new@x.com", "Inactive"),
    ])
    output = tmp_path / "report.xlsx"

    run_reconciliation(file_a, file_b, output)

    main = read_main(output)
    details = read_details(output)

    # One summary row per Customer ID in the main sheet
    assert len(main) == 2
    assert main.loc[1, "Status"] == "MATCHED"
    assert main.loc[10, "Status"] == "DUPLICATE"
    assert main.loc[10, "Duplicate Rows"] == "File A: 3, 4; File B: 3"
    assert "File A" in main.loc[10, "Changes"]
    assert "File B" not in main.loc[10, "Changes"]

    # Every occurrence, from both files, is listed once
    assert detail_occurrences(details, 10) == [
        ("File A", 3, "rui.old@x.com"),
        ("File A", 4, "rui.new@x.com"),
        ("File B", 3, "rui.new@x.com"),
    ]
    assert list(
        details[details["Customer ID"] == 10]["Occurrences in File"]
    ) == [2, 2, 1]


def test_summary_row_does_not_pick_one_of_several_different_values(
    make_excel, tmp_path
):

    file_a = make_excel("a.xlsx", [
        (10, "Rui", "rui.old@x.com", "Active"),
        (10, "Rui", "rui.new@x.com", "Inactive"),
    ])
    file_b = make_excel("b.xlsx", [
        (10, "Rui", "rui.new@x.com", "Inactive"),
    ])
    output = tmp_path / "report.xlsx"

    run_reconciliation(file_a, file_b, output)

    row = read_main(output).loc[10]

    # Same value in every occurrence: shown
    assert row["Name_A"] == "Rui"
    assert row["Name_B"] == "Rui"
    # Different values between the occurrences: not shown as if it were one
    assert row["Email_A"] == "(different values)"
    assert row["Status_A"] == "(different values)"
    # Single occurrence in File B: shown
    assert row["Email_B"] == "rui.new@x.com"
    assert row["Status_B"] == "Inactive"


# --- Duplicates only in File B ---

def test_duplicate_only_in_file_b_is_not_reported_as_new(
    make_excel, tmp_path
):

    file_a = make_excel("a.xlsx", [
        (1, "Ana", "ana@x.com", "Active"),
    ])
    file_b = make_excel("b.xlsx", [
        (1, "Ana", "ana@x.com", "Active"),
        (20, "Eva", "eva.1@x.com", "Active"),
        (20, "Eva", "eva.2@x.com", "Closed"),
    ])
    output = tmp_path / "report.xlsx"

    run_reconciliation(file_a, file_b, output)

    main = read_main(output)
    details = read_details(output)

    assert len(main) == 2
    assert main.loc[20, "Status"] == "DUPLICATE"
    assert main.loc[20, "Duplicate Rows"] == "File B: 3, 4"
    assert "File B" in main.loc[20, "Changes"]
    assert "File A" not in main.loc[20, "Changes"]

    # Not present in File A: the File A columns stay empty
    assert pd.isna(main.loc[20, "Name_A"])
    assert pd.isna(main.loc[20, "Email_A"])
    assert main.loc[20, "Email_B"] == "(different values)"

    assert detail_occurrences(details, 20) == [
        ("File B", 3, "eva.1@x.com"),
        ("File B", 4, "eva.2@x.com"),
    ]


# --- Duplicates in both files, with different values ---

def test_duplicate_in_both_files_has_no_cartesian_combinations(
    make_excel, tmp_path
):

    file_a = make_excel("a.xlsx", [
        (7, "Rui", "a1@x.com", "Active"),
        (7, "Rui", "a2@x.com", "Inactive"),
    ])
    file_b = make_excel("b.xlsx", [
        (7, "Rui", "b1@x.com", "Active"),
        (7, "Rui", "b2@x.com", "Closed"),
    ])
    output = tmp_path / "report.xlsx"

    run_reconciliation(file_a, file_b, output)

    main = read_main(output)
    details = read_details(output)

    # 2 x 2 occurrences are 4 real records, not 4 A/B "pairs"
    assert len(main) == 1
    assert main.loc[7, "Status"] == "DUPLICATE"
    assert main.loc[7, "Duplicate Rows"] == "File A: 2, 3; File B: 2, 3"
    assert main.loc[7, "Changes"].startswith(
        "Customer ID appears multiple times in File A and File B"
    )

    assert list(details.columns) == DETAILS_COLUMNS
    assert detail_occurrences(details, 7) == [
        ("File A", 2, "a1@x.com"),
        ("File A", 3, "a2@x.com"),
        ("File B", 2, "b1@x.com"),
        ("File B", 3, "b2@x.com"),
    ]
    assert list(details["Status"]) == [
        "Active", "Inactive", "Active", "Closed"
    ]


def test_every_occurrence_of_a_duplicate_id_is_in_the_details(
    make_excel, tmp_path
):

    file_a = make_excel("a.xlsx", [
        (1, "Ana", "ana@x.com", "Active"),
        (5, "Carl", "c1@x.com", "Active"),
        (2, "Maria", "m@x.com", "Active"),
        (5, "Carl", "c2@x.com", "Inactive"),
        (5, "Carlos", "c3@x.com", "Active"),
        (9, "Zoe", "z@x.com", "Active"),
    ])
    file_b = make_excel("b.xlsx", [
        (5, "Carl", "c4@x.com", "Active"),
        (2, "Maria", "m@x.com", "Active"),
        (2, "Maria", "m.other@x.com", "Active"),
        (9, "Zoe", "z@x.com", "Active"),
    ])
    output = tmp_path / "report.xlsx"

    run_reconciliation(file_a, file_b, output)

    df_a = pd.read_excel(file_a)
    df_b = pd.read_excel(file_b)

    duplicate_ids = set(
        df_a.loc[df_a["Customer ID"].duplicated(keep=False), "Customer ID"]
    ) | set(
        df_b.loc[df_b["Customer ID"].duplicated(keep=False), "Customer ID"]
    )

    assert duplicate_ids == {2, 5}

    # Every input row of a duplicated Customer ID, with its source and its
    # Excel row number
    expected = []

    for label, df in (("File A", df_a), ("File B", df_b)):
        for index, row in df.iterrows():
            if row["Customer ID"] in duplicate_ids:
                expected.append((
                    int(row["Customer ID"]),
                    label,
                    index + 2,
                    row["Name"],
                    row["Email"],
                    row["Status"],
                ))

    details = read_details(output)

    actual = [
        (
            int(row["Customer ID"]),
            row["Source File"],
            int(row["Excel Row"]),
            row["Name"],
            row["Email"],
            row["Status"],
        )
        for _, row in details.iterrows()
    ]

    assert len(actual) == len(expected) == 7
    assert sorted(actual) == sorted(expected)

    # Customer ID 9 is unique in each file, so it is not a duplicate
    assert 9 not in set(details["Customer ID"])


def test_duplicate_rows_lists_non_adjacent_excel_rows(make_excel, tmp_path):

    file_a = make_excel("a.xlsx", [
        (5, "Carl", "c1@x.com", "Active"),
        (6, "Dan", "d@x.com", "Active"),
        (5, "Carl", "c2@x.com", "Active"),
        (7, "Eve", "e@x.com", "Active"),
        (5, "Carl", "c3@x.com", "Active"),
    ])
    file_b = make_excel("b.xlsx", [
        (6, "Dan", "d@x.com", "Active"),
    ])
    output = tmp_path / "report.xlsx"

    run_reconciliation(file_a, file_b, output)

    main = read_main(output)

    assert main.loc[5, "Duplicate Rows"] == "File A: 2, 4, 6"
    assert len(read_details(output)) == 3


def test_identical_duplicates_are_still_duplicates(make_excel, tmp_path):

    file_a = make_excel("a.xlsx", [
        (5, "Carl", "carl@x.com", "Active"),
        (5, "Carl", "carl@x.com", "Active"),
    ])
    file_b = make_excel("b.xlsx", [
        (5, "Carl", "carl@x.com", "Active"),
    ])
    output = tmp_path / "report.xlsx"

    run_reconciliation(file_a, file_b, output)

    main = read_main(output)

    assert main.loc[5, "Status"] == "DUPLICATE"
    assert main.loc[5, "Email_A"] == "carl@x.com"
    assert len(read_details(output)) == 3


def test_report_with_only_duplicated_ids_is_generated(make_excel, tmp_path):

    file_a = make_excel("a.xlsx", [
        (1, "Ana", "a1@x.com", "Active"),
        (1, "Ana", "a2@x.com", "Active"),
    ])
    file_b = make_excel("b.xlsx", [
        (1, "Ana", "a1@x.com", "Active"),
    ])
    output = tmp_path / "report.xlsx"

    run_reconciliation(file_a, file_b, output)

    main = read_main(output)

    assert list(main.index) == [1]
    assert main.loc[1, "Status"] == "DUPLICATE"
    assert len(read_details(output)) == 3


# --- Unique IDs are not affected ---

def test_unique_ids_keep_their_status_and_have_no_duplicate_info(
    valid_files, tmp_path
):

    file_a, file_b = valid_files
    output = tmp_path / "report.xlsx"

    run_reconciliation(file_a, file_b, output)

    main = read_main(output)
    details = read_details(output)

    assert list(main.index) == [1, 2, 3, 4, 5]
    assert list(main["Status"]) == [
        "MATCHED", "CHANGED", "MISSING", "NEW", "DUPLICATE"
    ]

    for customer_id in (1, 2, 3, 4):
        assert pd.isna(main.loc[customer_id, "Duplicate Rows"])

    assert main.loc[2, "Changes"] == (
        "Email: maria@example.com -> maria.new@example.com"
    )

    # Only the duplicated Customer ID is in the details
    assert set(details["Customer ID"]) == {5}
    assert detail_occurrences(details, 5) == [
        ("File A", 5, "carl@example.com"),
        ("File A", 6, "carl@example.com"),
        ("File B", 5, "carl@example.com"),
    ]


def test_report_without_duplicates_has_no_details_sheet(
    make_excel, tmp_path
):

    file_a = make_excel("a.xlsx", [
        (1, "Ana", "ana@x.com", "Active"),
        (2, "Rui", "rui@x.com", "Active"),
    ])
    file_b = make_excel("b.xlsx", [
        (1, "Ana", "ana@x.com", "Active"),
        (3, "Eva", "eva@x.com", "Active"),
    ])
    output = tmp_path / "report.xlsx"

    run_reconciliation(file_a, file_b, output)

    workbook = load_workbook(output)

    assert workbook.sheetnames == ["Sheet1"]
    assert list(read_main(output)["Status"]) == [
        "MATCHED", "MISSING", "NEW"
    ]


# --- Excel file structure ---

def test_details_sheet_is_formatted_and_main_sheet_stays_first(
    valid_files, tmp_path
):

    file_a, file_b = valid_files
    output = tmp_path / "report.xlsx"

    run_reconciliation(file_a, file_b, output)

    workbook = load_workbook(output)

    assert workbook.sheetnames == ["Sheet1", DETAILS_SHEET]

    details = workbook[DETAILS_SHEET]

    assert [cell.value for cell in details[1]] == DETAILS_COLUMNS
    assert all(cell.font.bold for cell in details[1])
    assert details.freeze_panes == "A2"
    assert details.auto_filter.ref == details.dimensions
    assert details.max_row == 4

    # The main sheet keeps its formatting and its conditional formatting
    main = workbook["Sheet1"]

    assert main.freeze_panes == "A2"
    assert main.conditional_formatting
