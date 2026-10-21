from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule


DETAILS_SHEET_NAME = "Duplicate Details"

MAIN_COLUMN_WIDTHS = {
    "A": 14,
    "B": 20,
    "C": 20,
    "D": 28,
    "E": 28,
    "F": 18,
    "G": 18,
    "H": 14,
    "I": 60,
    "J": 30
}

DETAILS_COLUMN_WIDTHS = {
    "A": 14,
    "B": 14,
    "C": 12,
    "D": 20,
    "E": 28,
    "F": 14,
    "G": 20
}


def format_sheet(worksheet, column_widths):

    # Header formatting
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Freeze the first row
    worksheet.freeze_panes = "A2"

    # Enable filters
    worksheet.auto_filter.ref = worksheet.dimensions

    # Adjust column widths
    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    # Align the data
    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center")


def create_excel_report(result_df, output_file, duplicate_details=None):

    # Create the output folder if it does not exist yet
    Path(output_file).parent.mkdir(parents=True, exist_ok=True)

    # Replace missing values with empty cells
    result_df = result_df.fillna("")

    has_details = (
        duplicate_details is not None and not duplicate_details.empty
    )

    # Save the report to Excel. The Duplicate Details sheet is only created
    # when there are duplicated Customer IDs.
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:

        result_df.to_excel(writer, index=False)

        if has_details:
            duplicate_details.fillna("").to_excel(
                writer,
                sheet_name=DETAILS_SHEET_NAME,
                index=False
            )

    # Open the generated file to apply formatting
    workbook = load_workbook(output_file)
    worksheet = workbook.worksheets[0]

    format_sheet(worksheet, MAIN_COLUMN_WIDTHS)

    if has_details:
        format_sheet(workbook[DETAILS_SHEET_NAME], DETAILS_COLUMN_WIDTHS)

    # Conditional formatting for the Status column
    status_colors = {
        "MATCHED": "C6EFCE",
        "CHANGED": "FFEB9C",
        "MISSING": "FFC7CE",
        "NEW": "BDD7EE",
        "DUPLICATE": "F4CCCC"
    }

    for status, color in status_colors.items():

        fill = PatternFill(
            start_color=color,
            end_color=color,
            fill_type="solid"
        )

        worksheet.conditional_formatting.add(
            f"H2:H{worksheet.max_row}",
            FormulaRule(
                formula=[f'H2="{status}"'],
                fill=fill
            )
        )

    # Save again
    workbook.save(output_file)

    print(f"\nReport created successfully: {output_file}")
